"""Tests wiring contextual table styles into build_docx_report.py (issue #13).

Mirrors tools/test_table_styles_latex_wiring.py: `table_styles_enabled`
(default False) leaves the legacy generic renderer unchanged; only an
opted-in report activates directive-driven per-table selection, with no
silent fallback for an undirected table.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest
from docx.oxml.ns import qn

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))

import build_docx_report  # noqa: E402
from report_config import ReportConfig  # noqa: E402

DIRECTED_MD = """\
# Resultados

<!-- table-style: results-summary purpose=reference -->
| Name | Score |
| ---- | ----- |
| Ana  | 9     |
| Luis | 7     |
"""

UNDIRECTED_MD = """\
# Resultados

| Name | Score |
| ---- | ----- |
| Ana  | 9     |
"""


def _config(tmp_path: Path, extra: str = "") -> ReportConfig:
    folder = tmp_path / "report"
    folder.mkdir()
    (folder / "report.yml").write_text(
        "type: essay\nbackend: docx\noutput: docx\nmetadata:\n  title: Informe\n" + extra,
        encoding="utf-8",
    )
    (folder / "body.md").write_text(DIRECTED_MD, encoding="utf-8")
    from report_config import load_report_config

    return load_report_config(folder)


def _shd_fill(cell) -> str | None:
    shd = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
    return shd.get(qn("w:fill")) if shd is not None else None


def test_table_styles_disabled_by_default_uses_legacy_render_table(tmp_path: Path) -> None:
    config = _config(tmp_path)
    document, _warnings = build_docx_report.render_document(config)
    table = document.tables[-1]
    assert _shd_fill(table.rows[0].cells[0]) == "EAEAEA"  # legacy generic header shading


def test_table_styles_enabled_renders_selected_style(tmp_path: Path) -> None:
    config = _config(tmp_path, "table_styles:\n  enabled: true\n")
    document, _warnings = build_docx_report.render_document(config)
    table = document.tables[-1]
    # TAB-CL-01: same gray_shaded header fill as legacy, but via the new path.
    assert _shd_fill(table.rows[0].cells[0]) == "EAEAEA"


def test_table_styles_enabled_applies_teacher_override(tmp_path: Path) -> None:
    config = _config(
        tmp_path,
        "table_styles:\n  enabled: true\n  overrides:\n    results-summary: TAB-MN-03\n",
    )
    document, _warnings = build_docx_report.render_document(config)
    table = document.tables[-1]
    # TAB-MN-03: plain header, no shading.
    assert _shd_fill(table.rows[0].cells[0]) is None


def test_table_styles_enabled_blocks_on_missing_directive(tmp_path: Path) -> None:
    folder = tmp_path / "report"
    folder.mkdir()
    (folder / "report.yml").write_text(
        "type: essay\nbackend: docx\noutput: docx\nmetadata:\n  title: Informe\n"
        "table_styles:\n  enabled: true\n",
        encoding="utf-8",
    )
    (folder / "body.md").write_text(UNDIRECTED_MD, encoding="utf-8")
    from report_config import load_report_config

    config = load_report_config(folder)
    with pytest.raises(SystemExit, match="table-style"):
        build_docx_report.render_document(config)
