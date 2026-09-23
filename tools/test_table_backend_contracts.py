"""R9 `test_issue_13_backend_style_coherence`, per backend.

One selected style ID, rendered through a backend's token mapper, must
consistently represent borders, header, alignment, padding, density,
row rhythm, palette, indicators, caption and notes -- never a mix of
fragments from different styles. This file is parametrized over the seven
approved IDs per backend; harness commands select a backend by substring
(`-k latex`, `-k docx`, `-k html`), so test names carry that suffix.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from table_docx_tokens import render_styled_table_docx  # noqa: E402
from table_html_tokens import render_styled_table_html  # noqa: E402
from table_latex_tokens import render_styled_table_latex  # noqa: E402
from table_styles import APPROVED_STYLE_IDS, load_catalog  # noqa: E402

CATALOG = load_catalog()

HEADER = ["Item", "Status"]
ROWS = [["Widget A", "[[status:ok]]"], ["Widget B", "[[status:fail]]"]]


def _identity_convert_inline(text: str) -> str:
    return text


def _plain_fill_cell(cell, text: str, *, bold: bool, size_pt: float) -> None:
    """Minimal ``DocxRenderer._fill_cell`` stand-in: no citations/segments."""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = None if size_pt is None else run.font.size
    from docx.shared import Pt as _Pt

    run.font.size = _Pt(size_pt)


def _shd_fill(cell) -> str | None:
    shd = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
    return shd.get(qn("w:fill")) if shd is not None else None


def _cell_text(cell) -> str:
    return "".join(run.text for paragraph in cell.paragraphs for run in paragraph.runs)


@pytest.mark.parametrize("style_id", sorted(APPROVED_STYLE_IDS))
def test_issue_13_backend_style_coherence_latex(style_id: str) -> None:
    style = CATALOG.styles[style_id]
    lines = render_styled_table_latex(
        header=HEADER,
        rows=ROWS,
        style=style,
        status_indicators=CATALOG.status_indicators,
        convert_inline=_identity_convert_inline,
        emphasis_column=1 if style.tokens["row_rhythm"] == "column_emphasis" else None,
    )
    text = "\n".join(lines)

    # Identifies the selected style (evidence, never ambiguous).
    assert f"% table-style: {style_id}" in text

    tokens = style.tokens

    # Borders
    if tokens["borders"] == "full_grid":
        assert "|" in text
    elif tokens["borders"] in ("horizontal_only", "minimal"):
        assert "| >{" not in text and "|>{" not in text

    # Header shading
    if tokens["header"] == "gray_shaded":
        assert r"\rowcolor[gray]{0.92}" in text
    elif tokens["header"] == "dark_shaded":
        assert r"\rowcolor[gray]{0.25}" in text
        assert r"\color{white}" in text
    else:
        assert r"\rowcolor" not in text

    # Alignment
    align_marker = {"centered": r"\centering", "left": r"\raggedright", "numeric_right": r"\raggedleft"}
    assert align_marker[tokens["alignment"]] in text

    # Padding -> \tabcolsep value
    padding_pt = {"generous": "6pt", "standard": "3pt", "compact": "2pt", "minimal": "1pt"}
    assert rf"\setlength{{\tabcolsep}}{{{padding_pt[tokens['padding']]}}}" in text

    # Density -> font size
    font_size = {"low": r"\small", "medium": r"\footnotesize", "high": r"\scriptsize"}
    assert font_size[tokens["density"]] in text

    # Row rhythm
    if tokens["row_rhythm"] == "alternating":
        assert r"\rowcolors" in text
    elif tokens["row_rhythm"] == "column_emphasis":
        assert r"\cellcolor" in text
    else:
        assert r"\rowcolors" not in text

    # Indicators: never color alone -- symbol AND label always accompany color.
    if tokens["indicators"] == "symbol_color":
        assert "[[status:" not in text
        for value in ("ok", "fail"):
            indicator = CATALOG.status_indicators[value]
            assert indicator.symbol in text
            assert indicator.label in text
            assert indicator.color.lstrip("#") in text
    else:
        # A style with indicators=none still strips markers rather than
        # leaking raw authoring syntax into the PDF.
        assert "[[status:" not in text


def test_issue_13_backend_style_coherence_latex_caption_and_notes() -> None:
    style = CATALOG.styles["TAB-CL-01"]  # tokens: caption=above, notes=below
    lines = render_styled_table_latex(
        header=HEADER, rows=[["Widget A", "n/d"]], style=style,
        status_indicators=CATALOG.status_indicators,
        convert_inline=_identity_convert_inline,
        caption="Resultados", notes="Fuente: elaboración propia.",
    )
    text = "\n".join(lines)
    assert r"\caption{Resultados}" in text
    assert "Fuente: elaboración propia." in text

    style_inline_notes = CATALOG.styles["TAB-ES-06"]  # tokens: notes=inline
    lines = render_styled_table_latex(
        header=HEADER, rows=[["Widget A", "[[status:ok]]"]], style=style_inline_notes,
        status_indicators=CATALOG.status_indicators,
        convert_inline=_identity_convert_inline,
        notes="Ver leyenda.",
    )
    text = "\n".join(lines)
    assert r"\multicolumn" in text
    assert "Ver leyenda." in text


def test_issue_13_backend_style_coherence_latex_unknown_status_marker_blocks() -> None:
    style = CATALOG.styles["TAB-TC-02"]
    with pytest.raises(ValueError, match="status"):
        render_styled_table_latex(
            header=HEADER, rows=[["Widget A", "[[status:mystery]]"]], style=style,
            status_indicators=CATALOG.status_indicators,
            convert_inline=_identity_convert_inline,
        )


def _braces_balanced(segment: str) -> bool:
    return segment.count("{") == segment.count("}")


@pytest.mark.parametrize(
    "style_id",
    sorted(sid for sid in APPROVED_STYLE_IDS if CATALOG.styles[sid].tokens["header"] == "dark_shaded"),
)
def test_issue_13_backend_style_coherence_latex_dark_header_braces_stay_in_cell(style_id: str) -> None:
    """R3: a dark_shaded header group must not span `&` or the row terminator.

    A TeX group opened before the first cell and closed after `\\\\` makes
    `&` end that group early -- "Missing } inserted" / "Extra alignment
    tab" when compiling.
    """
    style = CATALOG.styles[style_id]
    lines = render_styled_table_latex(
        header=HEADER, rows=ROWS, style=style,
        status_indicators=CATALOG.status_indicators,
        convert_inline=_identity_convert_inline,
    )
    header_row = lines[lines.index(r"\rowcolor[gray]{0.25}") + 1]
    cells = header_row.rstrip().removesuffix(r" \\").split(" & ")
    assert all(_braces_balanced(cell) for cell in cells), header_row


def test_issue_13_backend_style_coherence_latex_column_emphasis_needs_index() -> None:
    style = CATALOG.styles["TAB-CE-05"]
    with pytest.raises(ValueError, match="emphasis_column"):
        render_styled_table_latex(
            header=HEADER, rows=ROWS, style=style,
            status_indicators=CATALOG.status_indicators,
            convert_inline=_identity_convert_inline, emphasis_column=None,
        )


# ---------------------------------------------------------------------------
# DOCX (T4)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("style_id", sorted(APPROVED_STYLE_IDS))
def test_issue_13_backend_style_coherence_docx(style_id: str) -> None:
    style = CATALOG.styles[style_id]
    document = Document()
    table = render_styled_table_docx(
        document=document,
        header=HEADER,
        rows=ROWS,
        style=style,
        status_indicators=CATALOG.status_indicators,
        fill_cell=_plain_fill_cell,
        emphasis_column=1 if style.tokens["row_rhythm"] == "column_emphasis" else None,
    )
    tokens = style.tokens

    header_fill = _shd_fill(table.rows[0].cells[0])
    if tokens["header"] == "gray_shaded":
        assert header_fill == "EAEAEA"
    elif tokens["header"] == "dark_shaded":
        assert header_fill == "404040"
    else:
        assert header_fill is None

    align_enum = {
        "centered": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "numeric_right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[tokens["alignment"]]
    assert table.rows[1].cells[0].paragraphs[0].alignment == align_enum

    size_pt = {"low": 11, "medium": 10, "high": 8}[tokens["density"]]
    body_run = table.rows[1].cells[0].paragraphs[0].runs[0]
    assert body_run.font.size.pt == size_pt

    borders_xml = table._tbl.tblPr.find(qn("w:tblBorders"))
    if tokens["borders"] == "full_grid":
        assert borders_xml.find(qn("w:insideV")).get(qn("w:val")) == "single"
    else:
        assert borders_xml.find(qn("w:insideV")).get(qn("w:val")) == "nil"

    if tokens["row_rhythm"] == "alternating":
        # Second data row (index 2 overall) is shaded for banding.
        assert _shd_fill(table.rows[2].cells[0]) is not None
    if tokens["row_rhythm"] == "column_emphasis":
        assert _shd_fill(table.rows[0].cells[1]) is not None
        assert _shd_fill(table.rows[1].cells[1]) is not None

    body_text = "".join(_cell_text(cell) for row in table.rows[1:] for cell in row.cells)
    assert "[[status:" not in body_text
    if tokens["indicators"] == "symbol_color":
        for value in ("ok", "fail"):
            indicator = CATALOG.status_indicators[value]
            assert indicator.symbol in body_text
            assert indicator.label in body_text


def test_issue_13_backend_style_coherence_docx_caption_and_notes() -> None:
    style = CATALOG.styles["TAB-CL-01"]
    document = Document()
    render_styled_table_docx(
        document=document, header=HEADER, rows=[["Widget A", "n/d"]], style=style,
        status_indicators=CATALOG.status_indicators, fill_cell=_plain_fill_cell,
        caption="Resultados", notes="Fuente: elaboración propia.",
    )
    paragraph_texts = [p.text for p in document.paragraphs]
    assert any("Resultados" in text for text in paragraph_texts)
    assert any("Fuente: elaboración propia." in text for text in paragraph_texts)


def test_issue_13_backend_style_coherence_docx_unknown_status_marker_blocks() -> None:
    style = CATALOG.styles["TAB-TC-02"]
    document = Document()
    with pytest.raises(ValueError, match="status"):
        render_styled_table_docx(
            document=document, header=HEADER, rows=[["Widget A", "[[status:mystery]]"]],
            style=style, status_indicators=CATALOG.status_indicators, fill_cell=_plain_fill_cell,
        )


# ---------------------------------------------------------------------------
# HTML (T5)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("style_id", sorted(APPROVED_STYLE_IDS))
def test_issue_13_backend_style_coherence_html(style_id: str) -> None:
    style = CATALOG.styles[style_id]
    html = render_styled_table_html(
        header=HEADER, rows=ROWS, style=style,
        status_indicators=CATALOG.status_indicators,
        escape_inline=_identity_convert_inline,
        emphasis_column=1 if style.tokens["row_rhythm"] == "column_emphasis" else None,
    )
    tokens = style.tokens

    assert f'data-table-style="{style_id}"' in html

    if tokens["header"] == "gray_shaded":
        assert "#ededed" in html
    elif tokens["header"] == "dark_shaded":
        assert "#404040" in html
        assert "#ffffff" in html

    align_css = {"centered": "center", "left": "left", "numeric_right": "right"}[tokens["alignment"]]
    assert f"text-align:{align_css}" in html

    if tokens["row_rhythm"] == "alternating":
        assert "#f2f2f2" in html
    if tokens["row_rhythm"] == "column_emphasis":
        assert "#d9d9d9" in html

    assert "[[status:" not in html
    if tokens["indicators"] == "symbol_color":
        for value in ("ok", "fail"):
            indicator = CATALOG.status_indicators[value]
            assert indicator.symbol in html
            assert indicator.label in html
            assert indicator.color.lower() in html.lower()


def test_issue_13_backend_style_coherence_html_caption_and_notes() -> None:
    style = CATALOG.styles["TAB-CL-01"]
    html = render_styled_table_html(
        header=HEADER, rows=[["Widget A", "n/d"]], style=style,
        status_indicators=CATALOG.status_indicators, escape_inline=_identity_convert_inline,
        caption="Resultados", notes="Fuente: elaboración propia.",
    )
    assert "<caption" in html and "Resultados" in html
    assert "Fuente: elaboración propia." in html


def test_issue_13_backend_style_coherence_html_unknown_status_marker_blocks() -> None:
    style = CATALOG.styles["TAB-TC-02"]
    with pytest.raises(ValueError, match="status"):
        render_styled_table_html(
            header=HEADER, rows=[["Widget A", "[[status:mystery]]"]], style=style,
            status_indicators=CATALOG.status_indicators, escape_inline=_identity_convert_inline,
        )


def test_issue_13_backend_style_coherence_html_column_emphasis_needs_index() -> None:
    style = CATALOG.styles["TAB-CE-05"]
    with pytest.raises(ValueError, match="emphasis_column"):
        render_styled_table_html(
            header=HEADER, rows=ROWS, style=style,
            status_indicators=CATALOG.status_indicators, escape_inline=_identity_convert_inline,
            emphasis_column=None,
        )


def test_issue_13_backend_style_coherence_docx_column_emphasis_needs_index() -> None:
    style = CATALOG.styles["TAB-CE-05"]
    document = Document()
    with pytest.raises(ValueError, match="emphasis_column"):
        render_styled_table_docx(
            document=document, header=HEADER, rows=ROWS, style=style,
            status_indicators=CATALOG.status_indicators, fill_cell=_plain_fill_cell,
            emphasis_column=None,
        )
