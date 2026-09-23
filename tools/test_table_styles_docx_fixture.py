"""End-to-end T4 harness: the real build_docx_report.py pipeline.

Mirrors the harness command from odd/tasks/contextual-table-styles.md
(`.venv/bin/python tools/build_docx_report.py
tests/fixtures/table_styles/sample-reports/docx`), but asserts against a
REOPENED document (``Document(path)``) rather than the in-memory object,
following this repository's own established DOCX test convention
(tools/test_build_docx_report.py) -- a DOCX is a zip of XML parts, and
only a round trip proves a style survived serialization.

Deviates from tasks #6424's suggested raw ``word/document.xml`` diff
against a golden file: python-docx's own XML serialization is not stable
across environments (attribute ordering, whitespace), so a byte-diff
golden would be fragile in a way the project's existing DOCX tests
deliberately avoid. Property assertions on the reopened document give the
same coverage without that fragility.

Named ``sample-reports/``, not ``reports/`` -- see the LaTeX fixture for
why (``.gitignore``'s unanchored ``reports/`` rule).
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))

import build_docx_report  # noqa: E402

FIXTURE = TOOLS.parent / "tests" / "fixtures" / "table_styles" / "sample-reports" / "docx"


def _shd_fill(cell) -> str | None:
    shd = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
    return shd.get(qn("w:fill")) if shd is not None else None


def _cell_text(cell) -> str:
    return "".join(run.text for paragraph in cell.paragraphs for run in paragraph.runs)


def test_fixture_report_renders_all_three_selected_styles(tmp_path: Path) -> None:
    folder = tmp_path / "docx-fixture"
    shutil.copytree(FIXTURE, folder)
    config = build_docx_report.build(folder)

    document = Document(str(config.docx_path))
    # tables[0] is the academic-route cover metadata table (Route A always
    # renders it); the three body tables follow in document order.
    tables = document.tables
    assert len(tables) == 4

    results_summary, method_comparison, status_matrix = tables[1:]

    # results-summary -> TAB-CL-01 (gray shaded header, full grid borders)
    assert _shd_fill(results_summary.rows[0].cells[0]) == "EAEAEA"

    # method-comparison -> TAB-CE-05 (column emphasis on the recommended method)
    assert _shd_fill(method_comparison.rows[0].cells[2]) is not None
    assert _shd_fill(method_comparison.rows[1].cells[2]) is not None

    # status-matrix -> TAB-ES-06 (status indicators, never color alone)
    body_text = "".join(_cell_text(cell) for row in status_matrix.rows[1:] for cell in row.cells)
    assert "[[status:" not in body_text
    assert "✓" in body_text and "OK" in body_text
    assert "✗" in body_text and "Falla" in body_text
