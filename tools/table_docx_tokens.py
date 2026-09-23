"""Map contextual table style tokens to DOCX/OOXML (issue #13, spec R9).

``render_styled_table_docx`` is the DOCX backend's half of "coherent
cross-backend rendering": mirrors ``tools/table_latex_tokens.py`` token for
token, so LaTeX and DOCX represent the same selected style consistently
even though the underlying markup is unrelated.

Kept free of any import from ``tools/build_docx_report.py`` to avoid a
circular import (that module already imports from ``build_latex_report.py``
and will import this one): cell-text rendering is injected as ``fill_cell``
(matching ``DocxRenderer._fill_cell``'s signature) rather than imported.
"""
from __future__ import annotations

from typing import TYPE_CHECKING, Callable

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

if TYPE_CHECKING:
    from docx.document import Document
    from docx.table import Table
    from table_styles import StatusIndicator, StyleDefinition

from table_directives import strip_status_markers
from table_styles import COLUMN_EMPHASIS_FILL_HEX, HEADER_FILL_HEX, ROW_ALTERNATING_FILL_HEX

_ALIGN_ENUM = {
    "centered": WD_ALIGN_PARAGRAPH.CENTER,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "numeric_right": WD_ALIGN_PARAGRAPH.RIGHT,
}
_SIZE_PT = {"low": 11, "medium": 10, "high": 8}
# Mirrors table_latex_tokens's \tabcolsep pt mapping (left/right); top/bottom
# is half that, for a visually comparable (not pixel-identical) cell margin.
_PADDING_LR_DXA = {"generous": 120, "standard": 60, "compact": 40, "minimal": 20}
_HEADER_FILL = HEADER_FILL_HEX

# OOXML (ECMA-376) requires `w:tblPr`/`w:tcPr` children in a fixed schema
# order -- Word repairs or silently drops a document whose properties are
# out of sequence. python-docx does not expose `tblBorders`/`tcBorders`/
# `shd`/`tcMar` as typed accessors (only a few CT_TblPr/CT_TcPr children
# are), so this module inserts them itself via the same
# `insert_element_before` mechanism python-docx's own generated accessors
# use, keyed on each element's schema successors (CT_TblPr/CT_TcPr
# `_tag_seq` in ``docx/oxml/table.py``). This makes insertion order-
# independent: whichever of `_shade_cell`/`_set_cell_margins`/
# `_underline_header_row` runs first, each new child lands before its own
# schema successors, never merely appended at the end.
_TBL_PR_SUCCESSORS = (
    "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange",
)

# CT_TcPr child order (ECMA-376), from `w:tcBorders` onward -- the single
# source of truth every "successors of <tag>" tuple below derives from, so a
# reader never has to count elements in a positional slice to know what it
# means (T4+T5 review R2-docx-successor-slices: `_TC_PR_SHD_SUCCESSORS`/
# `_TC_PR_TCMAR_SUCCESSORS` used to be unexplained `[1:]`/`[3:]` slices of
# `_TC_PR_TCBORDERS_SUCCESSORS`).
_TC_PR_SCHEMA_TAIL = (
    "w:tcBorders", "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
    "w:vAlign", "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange",
)


def _successors_of(tag: str) -> tuple[str, ...]:
    """Every ``_TC_PR_SCHEMA_TAIL`` tag that must come after ``tag`` in a CT_TcPr."""
    return _TC_PR_SCHEMA_TAIL[_TC_PR_SCHEMA_TAIL.index(tag) + 1:]


_TC_PR_TCBORDERS_SUCCESSORS = _successors_of("w:tcBorders")
_TC_PR_SHD_SUCCESSORS = _successors_of("w:shd")
_TC_PR_TCMAR_SUCCESSORS = _successors_of("w:tcMar")


def _shade_cell(cell, fill_hex: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill_hex)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_pr.insert_element_before(shading, *_TC_PR_SHD_SUCCESSORS)


def _set_cell_margins(cell, lr_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    # CT_TcMar child order: top, left(start), bottom, right(end).
    for side, value in (("top", lr_dxa // 2), ("left", lr_dxa), ("bottom", lr_dxa // 2), ("right", lr_dxa)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tc_pr.insert_element_before(margins, *_TC_PR_TCMAR_SUCCESSORS)


def _set_table_borders(table: "Table", borders_token: str) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    single = borders_token == "full_grid"
    horizontal = borders_token in ("full_grid", "horizontal_only")
    # CT_TblBorders child order: top, left, bottom, right, insideH, insideV.
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{side}")
        vertical = side in ("left", "right", "insideV")
        keep = single or (horizontal and not vertical)
        element.set(qn("w:val"), "single" if keep else "nil")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    tbl_pr.insert_element_before(borders, *_TBL_PR_SUCCESSORS)


def _underline_header_row(table: "Table") -> None:
    """One explicit rule under the header row -- used for `borders: minimal`
    ("sin rejilla ... una línea de cabecera"), independent of the (all-nil)
    table-wide borders. Called after the header row exists."""
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:color"), "000000")
        tc_borders.append(bottom)
        tc_pr.insert_element_before(tc_borders, *_TC_PR_TCBORDERS_SUCCESSORS)


def _indicator_runs(paragraph, value: str, status_indicators: "dict[str, StatusIndicator]") -> None:
    if value not in status_indicators:
        raise ValueError(
            f"unknown status marker [[status:{value}]] -- approved values are "
            f"{sorted(status_indicators)}"
        )
    indicator = status_indicators[value]
    symbol_run = paragraph.add_run(f" {indicator.symbol}")
    symbol_run.font.color.rgb = RGBColor.from_string(indicator.color.lstrip("#"))
    # Label stays default-colored text: symbol conveys color, label conveys
    # words -- meaning never depends on color alone.
    paragraph.add_run(f" {indicator.label}")


def render_styled_table_docx(
    document: "Document",
    header: list[str],
    rows: list[list[str]],
    style: "StyleDefinition",
    status_indicators: "dict[str, StatusIndicator]",
    fill_cell: Callable[..., None],
    emphasis_column: int | None = None,
    caption: str | None = None,
    notes: str | None = None,
) -> "Table":
    """Render one table under one coherent resolved style.

    ``fill_cell(cell, text, *, bold, size_pt)`` matches
    ``DocxRenderer._fill_cell``'s signature, injected so this module never
    imports ``build_docx_report.py``. Raises ``ValueError`` for the same
    cases ``render_styled_table_latex`` does: a ``column_emphasis`` style
    with no (or an out-of-range) ``emphasis_column``, or an unapproved
    ``[[status:<value>]]`` marker.
    """
    tokens = style.tokens
    columns = len(header)

    if tokens["row_rhythm"] == "column_emphasis" and emphasis_column is None:
        raise ValueError(f"{style.id}: column_emphasis row rhythm requires an emphasis_column index")
    if emphasis_column is not None and not (0 <= emphasis_column < columns):
        raise ValueError(f"{style.id}: emphasis_column {emphasis_column} is out of range for {columns} columns")

    if caption and tokens["caption"] == "above":
        paragraph = document.add_paragraph()
        run = paragraph.add_run(caption)
        run.italic = True

    table = document.add_table(rows=0, cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table, tokens["borders"])

    size_pt = _SIZE_PT[tokens["density"]]
    lr_dxa = _PADDING_LR_DXA[tokens["padding"]]
    align_enum = _ALIGN_ENUM[tokens["alignment"]]
    is_emphasis_style = tokens["row_rhythm"] == "column_emphasis"
    header_fill = _HEADER_FILL.get(tokens["header"])

    legend_values: set[str] = set()

    def _emit_row(values: list[str], *, is_header: bool) -> None:
        if len(values) > columns:
            raise ValueError(
                f"{style.id}: row has {len(values)} cells but the header defines "
                f"{columns} columns -- a body row can never be wider than the header"
            )
        cells = table.add_row().cells
        padded = values + [""] * (columns - len(values))
        for column_index, raw_text in enumerate(padded):
            cell = cells[column_index]
            clean, markers = strip_status_markers(raw_text)
            bold = is_header
            is_emphasis_cell = is_emphasis_style and column_index == emphasis_column
            if is_emphasis_cell:
                bold = True
            fill_cell(cell, clean, bold=bold, size_pt=size_pt)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else align_enum
            _set_cell_margins(cell, lr_dxa)
            for value in markers:
                _indicator_runs(cell.paragraphs[0], value, status_indicators)
                legend_values.add(value)
            if is_header and header_fill:
                if tokens["header"] == "dark_shaded":
                    for run in cell.paragraphs[0].runs:
                        run.font.color.rgb = RGBColor.from_string("FFFFFF")
                _shade_cell(cell, header_fill)
            elif is_emphasis_cell:
                _shade_cell(cell, COLUMN_EMPHASIS_FILL_HEX)

    _emit_row(header, is_header=True)
    if tokens["borders"] == "minimal":
        _underline_header_row(table)
    for row_index, row in enumerate(rows):
        _emit_row(row, is_header=False)
        if tokens["row_rhythm"] == "alternating" and row_index % 2 == 1:
            for cell in table.rows[-1].cells:
                _shade_cell(cell, ROW_ALTERNATING_FILL_HEX)

    if notes and tokens["notes"] == "inline":
        note_cells = table.add_row().cells
        merged = note_cells[0]
        for extra in note_cells[1:]:
            merged = merged.merge(extra)
        run = merged.paragraphs[0].add_run(notes)
        run.italic = True

    if caption and tokens["caption"] == "below":
        paragraph = document.add_paragraph()
        run = paragraph.add_run(caption)
        run.italic = True
    if notes and tokens["notes"] == "below":
        paragraph = document.add_paragraph()
        run = paragraph.add_run(notes)
        run.italic = True
    if legend_values:
        legend_text = "  ".join(
            f"{status_indicators[value].symbol} {status_indicators[value].label}"
            for value in sorted(legend_values)
        )
        paragraph = document.add_paragraph()
        run = paragraph.add_run(legend_text)
        run.font.size = Pt(9)

    return table
